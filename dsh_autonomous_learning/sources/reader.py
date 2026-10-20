"""资料读取器。

支持：txt/md/html/pdf/docx/常见代码文件，以及 http(s) 网页。
不支持的类型必须明确报错（含扩展名与建议），不允许静默返回空内容。

安全约束：
- 本地文件受 PermissionManager 的授权根目录约束；
- 网络读取受 ``security.allow_network``、域名白名单与体积上限约束；
- 不绕过登录/付费墙；遇到明显需要登录的站点直接记录失败原因；
- 记录来源 URL/路径、访问时间与内容哈希，便于溯源。
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

from ..config import SecurityConfig
from ..core.models import utcnow_iso
from ..security.permissions import PermissionManager
from . import metadata as meta
from . import parser as parsing

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".xml"}
CODE_EXTENSIONS = meta.CODE_EXTENSIONS
HTML_EXTENSIONS = {".html", ".htm"}
DOC_TYPES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
}


class ReadError(RuntimeError):
    """读取失败（权限、格式、网络、体积）。"""


@dataclass
class ReadResult:
    ok: bool
    location: str = ""
    title: str = ""
    text: str = ""
    doc_type: str = "text"
    source_type: str = "unknown"
    content_hash: str = ""
    bytes_read: int = 0
    error: str = ""
    warnings: List[str] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "ok": self.ok,
            "location": self.location,
            "title": self.title,
            "doc_type": self.doc_type,
            "source_type": self.source_type,
            "content_hash": self.content_hash,
            "bytes_read": self.bytes_read,
            "char_count": len(self.text),
            "error": self.error,
            "warnings": list(self.warnings),
            "metadata": dict(self.metadata),
            "preview": self.text[:300],
        }


class SourceReader:
    def __init__(self, security: SecurityConfig, permissions: PermissionManager) -> None:
        self.security = security
        self.permissions = permissions

    # ------------------------------------------------------------------ #
    def read(self, location: str, *, project_id: str = "", confirm_token: str = "") -> ReadResult:
        if not location:
            return ReadResult(ok=False, error="location 不能为空")
        if location.startswith(("http://", "https://")):
            return self._read_http(location, confirm_token=confirm_token, project_id=project_id)
        return self._read_local(location, confirm_token=confirm_token, project_id=project_id)

    # ------------------------------------------------------------------ #
    def _read_local(self, location: str, *, confirm_token: str, project_id: str) -> ReadResult:
        path = Path(location).expanduser()
        if not path.exists():
            return ReadResult(ok=False, location=str(path), error=f"文件不存在: {path}")
        if path.is_dir():
            return ReadResult(ok=False, location=str(path), error="目标是目录，请使用 discover_sources 扫描目录")
        decision = self.permissions.decide(
            action="read_source", risk="read", target=str(path), project_id=project_id, confirmed_token=confirm_token
        )
        if not decision.allowed:
            return ReadResult(ok=False, location=str(path), error=f"权限拒绝: {decision.reason}")
        if decision.needs_confirmation and decision.allowed:
            return ReadResult(
                ok=False,
                location=str(path),
                error=(
                    "该路径需要用户确认。请先在 DSH 审批界面批准，或把路径加入 "
                    "security.allowed_roots / 复制到 data/corpus 后重试。"
                ),
            )

        size = path.stat().st_size
        if size > self.security.max_download_bytes:
            return ReadResult(
                ok=False,
                location=str(path),
                error=f"文件过大（{size} 字节 > 上限 {self.security.max_download_bytes}），已拒绝读取",
            )

        suffix = path.suffix.lower()
        doc_type = DOC_TYPES.get(suffix, "code" if suffix in CODE_EXTENSIONS else "text")
        warnings: List[str] = []
        try:
            if doc_type == "pdf":
                text = self._read_pdf(path)
            elif doc_type == "docx":
                text = self._read_docx(path)
            elif doc_type == "html":
                raw = self._decode_bytes(path.read_bytes())
                title, text = parsing.html_to_text(raw)
                return self._finish(path, title, text, "html", size, warnings)
            elif suffix in TEXT_EXTENSIONS or suffix in CODE_EXTENSIONS or suffix in HTML_EXTENSIONS:
                text = self._decode_bytes(path.read_bytes())
                if doc_type == "markdown":
                    text = parsing.markdown_to_text(text)
            else:
                return ReadResult(
                    ok=False,
                    location=str(path),
                    error=(
                        f"不支持的文件类型: {suffix or '(无扩展名)'}。"
                        "支持 txt/md/rst/html/pdf/docx 及常见代码文件；"
                        "如需其它格式，请先转换成纯文本再导入。"
                    ),
                )
        except ReadError as exc:
            return ReadResult(ok=False, location=str(path), error=str(exc))
        except Exception as exc:  # noqa: BLE001
            return ReadResult(ok=False, location=str(path), error=f"读取失败: {exc}")

        title = path.stem
        first_heading = re.search(r"^\s*#\s+(.+)$", text, re.M)
        if first_heading:
            title = first_heading.group(1).strip()
        return self._finish(path, title, text, doc_type, size, warnings)

    def _finish(
        self, path: Path, title: str, text: str, doc_type: str, size: int, warnings: List[str]
    ) -> ReadResult:
        text = (text or "").strip()
        if not text:
            return ReadResult(ok=False, location=str(path), error="解析结果为空：文件可能是扫描件或加密文档", doc_type=doc_type)
        digest = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()
        location = str(path)
        info = meta.describe_source(location, title=title, doc_type=doc_type)
        return ReadResult(
            ok=True,
            location=location,
            title=title or path.name,
            text=text,
            doc_type=doc_type,
            source_type=info["source_type"],
            content_hash=digest,
            bytes_read=size,
            warnings=warnings,
            metadata=info,
        )

    @staticmethod
    def _decode_bytes(raw: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    @staticmethod
    def _read_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise ReadError("解析 PDF 需要 pypdf，请先安装后再试") from exc
        try:
            reader = PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages):
                content = page.extract_text() or ""
                if content.strip():
                    pages.append(f"\n\n<!-- page {i + 1} -->\n{content}")
            text = "".join(pages)
        except Exception as exc:  # noqa: BLE001
            raise ReadError(f"PDF 解析失败: {exc}") from exc
        if not text.strip():
            raise ReadError("PDF 未提取到文本（可能是扫描件，需要 OCR，本插件不做 OCR）")
        return text

    @staticmethod
    def _read_docx(path: Path) -> str:
        try:
            import docx  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise ReadError("解析 DOCX 需要 python-docx，请先安装后再试") from exc
        try:
            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as exc:  # noqa: BLE001
            raise ReadError(f"DOCX 解析失败: {exc}") from exc

    # ------------------------------------------------------------------ #
    def _read_http(self, url: str, *, confirm_token: str, project_id: str) -> ReadResult:
        if meta.is_paywalled_or_login(url):
            return ReadResult(
                ok=False,
                location=url,
                error="该站点需要登录或有访问限制，插件不绕过登录/付费墙，已记录失败原因",
            )
        decision = self.permissions.decide(
            action="fetch_url", risk="external", target=url, project_id=project_id, confirmed_token=confirm_token
        )
        if not decision.allowed:
            return ReadResult(ok=False, location=url, error=f"网络访问被拒绝: {decision.reason}")
        if decision.needs_confirmation:
            return ReadResult(
                ok=False,
                location=url,
                error=(
                    "assisted 模式下外部抓取需要确认。请先调用 request_confirmation 获取 "
                    "confirm_token 后重试。"
                ),
            )

        host = urlparse(url).netloc.lower()
        if self.security.allowed_domains and not any(
            host == d or host.endswith("." + d.lstrip(".")) for d in self.security.allowed_domains
        ):
            return ReadResult(ok=False, location=url, error=f"域名 {host} 不在允许列表内")

        import httpx

        headers = {
            "User-Agent": "DSH-Autonomous-Learning/1.0 (+local research agent)",
            "Accept": "text/html,application/xhtml+xml,text/plain,application/pdf;q=0.8,*/*;q=0.5",
        }
        try:
            with httpx.Client(timeout=30, follow_redirects=True, headers=headers) as client:
                with client.stream("GET", url) as resp:
                    if resp.status_code >= 400:
                        return ReadResult(
                            ok=False, location=url, error=f"HTTP {resp.status_code}: 抓取失败"
                        )
                    content_type = (resp.headers.get("content-type") or "").lower()
                    body = bytearray()
                    for chunk in resp.iter_bytes():
                        body.extend(chunk)
                        if len(body) > self.security.max_download_bytes:
                            return ReadResult(
                                ok=False,
                                location=url,
                                error=f"响应体超过上限 {self.security.max_download_bytes} 字节，已中止",
                            )
                    final_url = str(resp.url)
        except Exception as exc:  # noqa: BLE001
            return ReadResult(ok=False, location=url, error=f"网络请求失败: {exc}")

        raw = bytes(body)
        if "pdf" in content_type:
            tmp = Path(self.permissions.data_dir or ".") / "documents" / "_tmp_fetch.pdf"
            tmp.parent.mkdir(parents=True, exist_ok=True)
            tmp.write_bytes(raw)
            try:
                text = self._read_pdf(tmp)
                title = url.rsplit("/", 1)[-1] or "PDF 文档"
                return self._finish_http(url, final_url, title, text, "pdf", len(raw))
            finally:
                tmp.unlink(missing_ok=True)
        text_raw = self._decode_bytes(raw)
        if "html" in content_type or text_raw.lstrip().lower().startswith(("<!doctype", "<html")):
            title, text = parsing.html_to_text(text_raw)
            return self._finish_http(url, final_url, title, text, "html", len(raw))
        return self._finish_http(url, final_url, url.rsplit("/", 1)[-1], text_raw, "text", len(raw))

    def _finish_http(self, url: str, final_url: str, title: str, text: str, doc_type: str, size: int) -> ReadResult:
        text = (text or "").strip()
        if not text:
            return ReadResult(ok=False, location=url, error="网页正文为空（可能需要 JS 渲染或存在访问限制）")
        digest = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()
        info = meta.describe_source(final_url, title=title, doc_type=doc_type)
        return ReadResult(
            ok=True,
            location=final_url,
            title=title or final_url,
            text=text,
            doc_type=doc_type,
            source_type=info["source_type"],
            content_hash=digest,
            bytes_read=size,
            metadata={**info, "requested_url": url, "accessed_at": utcnow_iso()},
        )


__all__ = ["DOC_TYPES", "ReadError", "ReadResult", "SourceReader", "TEXT_EXTENSIONS"]
